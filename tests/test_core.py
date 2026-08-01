from __future__ import annotations

import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from govmatch_core import (
    MAX_FILE_BYTES,
    Profile,
    check_fresh_graduate,
    check_household,
    check_major,
    dataframe_to_excel_bytes,
    evaluate_job,
    load_tabular_file,
    mapping_conflicts,
)


def profile(**overrides: object) -> Profile:
    values: dict[str, object] = {
        "education": "本科",
        "degree": "学士",
        "majors": ["法学"],
        "major_categories": ["法学类"],
        "political": "群众",
        "grassroots_years": 0,
        "service_project": "无",
        "fresh_graduate": False,
        "fresh_status": "非应届毕业生",
        "gender": "女",
        "household": "浙江省嘉兴市",
    }
    values.update(overrides)
    return Profile(**values)  # type: ignore[arg-type]


@pytest.mark.parametrize("requirement", ["非法学类", "法学类（不含法学）", "法学除外"])
def test_major_exclusion_beats_substring_match(requirement: str) -> None:
    result = check_major(requirement, profile())
    assert result.status == "fail"
    assert "排除" in result.reason


def test_major_positive_match_still_passes() -> None:
    assert check_major("法学类、政治学类", profile()).status == "pass"


def test_city_requirement_does_not_fall_back_to_same_province() -> None:
    result = check_household("浙江省杭州市户籍", "浙江省嘉兴市")
    assert result.status == "fail"


def test_province_requirement_accepts_city_within_province() -> None:
    assert check_household("浙江省户籍", "浙江省嘉兴市").status == "pass"


def test_unlimited_fresh_wording_is_not_treated_as_fresh_only() -> None:
    result = check_fresh_graduate("不限应届毕业生", False, "非应届毕业生")
    assert result.status == "pass"


def test_specific_graduation_year_requires_same_year() -> None:
    assert check_fresh_graduate("限2026届高校毕业生", True, "2026届高校毕业生").status == "pass"
    assert check_fresh_graduate("限2026届高校毕业生", True, "2025届高校毕业生").status == "fail"


def test_missing_columns_reduce_coverage_and_prevent_high_confidence() -> None:
    row = pd.Series({"学历": "本科及以上", "专业": "法学类"})
    result = evaluate_job(row, {"education": "学历", "major": "专业"}, profile())
    assert result.coverage_score == 20
    assert result.overall == "存在待核对项"
    assert result.checked_fields == 2
    assert "学位要求" in result.missing_fields


def test_export_neutralizes_formula_cells() -> None:
    source = pd.DataFrame({"职位名称": ['=HYPERLINK("https://example.test")', "+SUM(1,2)"]})
    workbook = load_workbook(io.BytesIO(dataframe_to_excel_bytes(source, profile())), data_only=False)
    sheet = workbook["岗位匹配结果"]
    assert sheet["A2"].value.startswith("'=")
    assert sheet["A3"].value.startswith("'+")
    assert "规则与来源" in workbook.sheetnames


def test_upload_size_limit_is_enforced_before_parsing() -> None:
    with pytest.raises(ValueError, match="20 MB"):
        load_tabular_file(b"x" * (MAX_FILE_BYTES + 1), "jobs.csv")


def test_low_confidence_excel_header_is_rejected() -> None:
    buffer = io.BytesIO()
    pd.DataFrame([["甲", "乙"], [1, 2]]).to_excel(buffer, index=False, header=False)
    with pytest.raises(ValueError, match="可靠识别表头"):
        load_tabular_file(buffer.getvalue(), "unknown.xlsx")


def test_mapping_conflict_is_reported() -> None:
    assert mapping_conflicts({"major": "条件", "education": "条件", "degree": None}) == {
        "条件": ["major", "education"]
    }
